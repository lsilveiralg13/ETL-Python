import base64
import asyncio
import io
import json
import unittest
import uuid
from contextlib import asynccontextmanager
from types import SimpleNamespace
from unittest.mock import patch
import zipfile

from qdrant_client import QdrantClient, models

from core.rag_guard import buscar_conhecimento_rag
from utils.meu_agente_dados import attachments, feedback, main, servidor_mcp


class CacheTests(unittest.TestCase):
    def setUp(self):
        servidor_mcp.CACHE_MEMORIA.clear()

    def tearDown(self):
        servidor_mcp.CACHE_MEMORIA.clear()

    def test_cache_tem_limite_e_remove_menos_recentemente_usada(self):
        chamadas = []

        @servidor_mcp.com_cache()
        def buscar(valor):
            chamadas.append(valor)
            return valor

        with patch.object(servidor_mcp, "CACHE_MAX_ENTRADAS", 2):
            self.assertEqual(buscar("a"), "a")
            self.assertEqual(buscar("b"), "b")
            self.assertIn("Resposta obtida do cache", buscar("a"))
            self.assertEqual(buscar("c"), "c")
            self.assertEqual(len(servidor_mcp.CACHE_MEMORIA), 2)
            self.assertEqual(buscar("b"), "b")

        self.assertEqual(chamadas, ["a", "b", "c", "b"])

    def test_cache_respeita_expiracao(self):
        chamadas = []

        @servidor_mcp.com_cache(ttl_segundos=0)
        def consultar():
            chamadas.append(None)
            return len(chamadas)

        self.assertEqual(consultar(), 1)
        self.assertEqual(consultar(), 2)


class AttachmentTests(unittest.TestCase):
    def test_decodificar_anexo_aceita_base64_valido(self):
        anexo = {
            "nome": "dados.txt",
            "mime_type": "text/plain",
            "conteudo_b64": base64.b64encode(b"dados").decode("ascii"),
        }

        conteudo, mime_type = main.decodificar_anexo(anexo)

        self.assertEqual(conteudo, b"dados")
        self.assertEqual(mime_type, "text/plain")

    def test_decodificar_anexo_rejeita_texto_codificado_acima_do_limite(self):
        anexo = {
            "nome": "grande.txt",
            "conteudo_b64": "A" * 5,
        }

        with patch.object(main, "MAX_ANEXO_ENCODED_CHARS", 4):
            with self.assertRaisesRegex(ValueError, "excede o limite"):
                main.decodificar_anexo(anexo)

    def test_chat_request_rejeita_tamanho_codificado_total_acima_do_limite(self):
        with patch.object(main, "MAX_TOTAL_ANEXOS_BYTES", 3):
            request = main.ChatRequest(
                prompt="Analise",
                anexos=[
                    main.Anexo(nome="a.txt", conteudo_b64="YQ=="),
                    main.Anexo(nome="b.txt", conteudo_b64="Yg=="),
                ],
            )
            self.assertEqual(len(request.anexos), 2)

            with self.assertRaisesRegex(ValueError, "total dos anexos"):
                main.ChatRequest(
                    prompt="Analise",
                    anexos=[main.Anexo(nome="grande.txt", conteudo_b64="A" * 29)],
                )

    def test_codigo_e_documentos_office_sao_enviados_como_texto_extraido(self):
        codigo = SimpleNamespace(
            name="consulta.sql",
            size=0,
            seek=lambda *_args: None,
            read=lambda _limit: b"SELECT 1 FROM dual;",
        )
        payload = attachments.preparar_anexo(codigo)
        conteudo = base64.b64decode(payload["conteudo_b64"]).decode("utf-8")
        self.assertEqual(payload["mime_type"], "text/plain")
        self.assertIn("SELECT 1 FROM dual;", conteudo)

        docx_bytes = io.BytesIO()
        namespace = (
            "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        )
        document = (
            f'<w:document xmlns:w="{namespace}"><w:body><w:p>'
            "<w:r><w:t>Conteúdo do documento</w:t></w:r></w:p></w:body></w:document>"
        )
        with zipfile.ZipFile(docx_bytes, "w") as archive:
            archive.writestr("word/document.xml", document)
        docx = SimpleNamespace(
            name="manual.docx",
            seek=lambda *_args: None,
            read=lambda _limit: docx_bytes.getvalue(),
        )
        self.assertIn("Conteúdo do documento", attachments.extrair_texto_upload(docx))

    def test_anexos_multimodais_preservam_mime_e_formatos_desconhecidos_sao_negados(self):
        imagem = SimpleNamespace(
            name="captura.webp",
            seek=lambda *_args: None,
            read=lambda _limit: b"imagem",
        )
        self.assertEqual(
            attachments.preparar_anexo(imagem)["mime_type"],
            "image/webp",
        )
        executavel = SimpleNamespace(
            name="programa.exe",
            seek=lambda *_args: None,
            read=lambda _limit: b"MZ",
        )
        with self.assertRaisesRegex(ValueError, "não suportado"):
            attachments.preparar_anexo(executavel)

    def test_planilhas_apresentacoes_e_limite_de_paginas_pdf(self):
        from openpyxl import Workbook
        from pptx import Presentation
        from pypdf import PdfWriter

        workbook = Workbook()
        workbook.active.append(["produto", "quantidade"])
        workbook.active.append(["A", 3])
        xlsx_bytes = io.BytesIO()
        workbook.save(xlsx_bytes)
        planilha = SimpleNamespace(
            name="dados.xlsx",
            seek=lambda *_args: None,
            read=lambda _limit: xlsx_bytes.getvalue(),
        )
        self.assertIn("produto\tquantidade", attachments.extrair_texto_upload(planilha))

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 1000, 500).text = "Resumo executivo"
        pptx_bytes = io.BytesIO()
        presentation.save(pptx_bytes)
        apresentacao = SimpleNamespace(
            name="resumo.pptx",
            seek=lambda *_args: None,
            read=lambda _limit: pptx_bytes.getvalue(),
        )
        self.assertIn("Resumo executivo", attachments.extrair_texto_upload(apresentacao))

        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        pdf_bytes = io.BytesIO()
        writer.write(pdf_bytes)
        pdf = SimpleNamespace(
            name="pagina.pdf",
            seek=lambda *_args: None,
            read=lambda _limit: pdf_bytes.getvalue(),
        )
        with patch.object(attachments, "MAX_PDF_PAGES", 0):
            with self.assertRaisesRegex(ValueError, "limite de 0 páginas"):
                attachments.preparar_anexo(pdf)


class AgentRuntimeTests(unittest.TestCase):
    def test_oracle_dialeto_e_forcado_na_chamada_de_validacao_sql(self):
        request = main.ChatRequest(prompt="Use Oracle", dialeto_sql="oracle")
        self.assertEqual(request.dialeto_sql, "Oracle")
        argumentos = main.preparar_argumentos_ferramenta(
            "validar_e_executar_sql",
            {"query": "SELECT 1 FROM dual", "dialecto": "mysql"},
            request.dialeto_sql,
        )
        self.assertEqual(argumentos["dialecto"], "oracle")
        with self.assertRaises(ValueError):
            main.ChatRequest(prompt="SQL", dialeto_sql="PL/SQL")
        self.assertEqual(
            main.normalizar_mime_type("application/octet-stream", "dados.csv"),
            "text/csv",
        )

    def test_subprocesso_mcp_recebe_apenas_variaveis_permitidas(self):
        with patch.dict(
            "os.environ",
            {
                "GEMINI_API_KEY": "gemini-secret",
                "QDRANT_API_KEY": "qdrant-secret",
                "DATABASE_PASSWORD": "database-secret",
                "PYTHONPATH": "caminho-nao-confiavel",
            },
            clear=True,
        ):
            ambiente = main.criar_env_subprocesso()

        self.assertNotIn("GEMINI_API_KEY", ambiente)
        self.assertNotIn("DATABASE_PASSWORD", ambiente)
        self.assertNotIn("caminho-nao-confiavel", ambiente["PYTHONPATH"])
        self.assertEqual(ambiente["QDRANT_API_KEY"], "qdrant-secret")

    def test_streaming_aggregates_text_and_function_call(self):
        class FakeModels:
            async def generate_content_stream(self, **_kwargs):
                yield SimpleNamespace(
                    candidates=[
                        SimpleNamespace(
                            content=main.types.Content(
                                role="model",
                                parts=[main.types.Part.from_text(text="Verificando ")],
                            )
                        )
                    ]
                )
                yield SimpleNamespace(
                    candidates=[
                        SimpleNamespace(
                            content=main.types.Content(
                                role="model",
                                parts=[
                                    main.types.Part.from_text(text="dados."),
                                    main.types.Part(
                                        function_call=main.types.FunctionCall(
                                            name="calcular_indicador_otif",
                                            args={"unidade": "Contagem"},
                                        ),
                                        thought_signature="c2ln",
                                    ),
                                ],
                            )
                        )
                    ]
                )

        class FakeClient:
            aio = SimpleNamespace(models=FakeModels())

        eventos = []

        async def registrar(evento):
            eventos.append(evento)

        resposta = asyncio.run(
            main.gerar_conteudo_streaming(
                FakeClient(),
                "modelo-teste",
                [],
                None,
                registrar,
            )
        )

        self.assertEqual(resposta.text, "Verificando dados.")
        self.assertEqual(resposta.function_calls[0].name, "calcular_indicador_otif")
        self.assertEqual(
            dict(resposta.function_calls[0].args),
            {"unidade": "Contagem"},
        )
        self.assertEqual(
            resposta.candidates[0].content.parts[1].thought_signature,
            b"sig",
        )
        self.assertEqual("".join(e["text"] for e in eventos), resposta.text)

    def test_streaming_aguarda_sdk_que_retorna_coroutine_do_stream(self):
        class FakeModels:
            async def generate_content_stream(self, **_kwargs):
                async def chunks():
                    yield SimpleNamespace(
                        candidates=[
                            SimpleNamespace(
                                content=main.types.Content(
                                    role="model",
                                    parts=[
                                        main.types.Part.from_text(
                                            text="Stream aguardado corretamente."
                                        )
                                    ],
                                )
                            )
                        ]
                    )

                return chunks()

        resposta = asyncio.run(
            main.gerar_conteudo_streaming(
                SimpleNamespace(aio=SimpleNamespace(models=FakeModels())),
                "modelo-teste",
                [],
                None,
            )
        )
        self.assertEqual(resposta.text, "Stream aguardado corretamente.")

    def test_fallback_tenta_proximo_modelo_apos_erro_no_stream(self):
        class FakeModels:
            async def generate_content_stream(self, *, model, **_kwargs):
                if model == "indisponivel":
                    raise RuntimeError("serviço temporariamente indisponível")
                yield SimpleNamespace(
                    candidates=[
                        SimpleNamespace(
                            content=main.types.Content(
                                role="model",
                                parts=[main.types.Part.from_text(text="Resposta alternativa.")],
                            )
                        )
                    ]
                )

        client = SimpleNamespace(aio=SimpleNamespace(models=FakeModels()))
        eventos = []

        async def registrar(evento):
            eventos.append(evento)

        resposta, modelo = asyncio.run(
            main.gerar_resposta_com_fallback(
                client,
                ["indisponivel", "alternativo"],
                [],
                None,
                registrar,
            )
        )

        self.assertEqual(modelo, "alternativo")
        self.assertEqual(resposta.text, "Resposta alternativa.")
        self.assertIn("reset", [evento["type"] for evento in eventos])
        self.assertEqual(
            [evento["name"] for evento in eventos if evento["type"] == "model"],
            ["indisponivel", "alternativo"],
        )

    def test_fallback_nao_repete_erro_deterministico_de_thought_signature(self):
        modelos_consultados = []

        class FakeModels:
            async def generate_content_stream(self, *, model, **_kwargs):
                modelos_consultados.append(model)
                raise RuntimeError(
                    "400 INVALID_ARGUMENT: thought_signature ausente"
                )

        client = SimpleNamespace(aio=SimpleNamespace(models=FakeModels()))

        with self.assertRaisesRegex(RuntimeError, "thought_signature"):
            asyncio.run(
                main.gerar_resposta_com_fallback(
                    client,
                    ["modelo-inicial", "modelo-alternativo"],
                    [],
                    None,
                )
            )

        self.assertEqual(modelos_consultados, ["modelo-inicial"])

    def test_fallback_explica_cota_esgotada_em_todos_os_modelos(self):
        modelos_consultados = []

        class FakeModels:
            async def generate_content_stream(self, *, model, **_kwargs):
                modelos_consultados.append(model)
                raise RuntimeError(
                    "429 RESOURCE_EXHAUSTED: quota exceeded"
                )

        client = SimpleNamespace(aio=SimpleNamespace(models=FakeModels()))

        with self.assertRaisesRegex(RuntimeError, "cota do Gemini está esgotada"):
            asyncio.run(
                main.gerar_resposta_com_fallback(
                    client,
                    ["modelo-a", "modelo-b"],
                    [],
                    None,
                )
            )

        self.assertEqual(modelos_consultados, ["modelo-a", "modelo-b"])

    def test_runtime_mantem_uma_sessao_mcp_ate_o_encerramento(self):
        inicializacoes = []
        encerramentos = []

        class FakeSession:
            async def __aenter__(self):
                return self

            async def __aexit__(self, *_args):
                encerramentos.append("sessao")

            async def initialize(self):
                inicializacoes.append("initialize")

            async def list_tools(self):
                return SimpleNamespace(tools=["ferramenta"])

            async def send_ping(self):
                return None

        @asynccontextmanager
        async def fake_stdio(_params):
            yield object(), object()
            encerramentos.append("stdio")

        async def executar():
            runtime = main.MCPRuntime()
            with (
                patch.object(main, "stdio_client", fake_stdio),
                patch.object(main, "ClientSession", return_value=FakeSession()),
            ):
                await runtime.start()
                await runtime.start()
                self.assertEqual(runtime.tools, ["ferramenta"])
                self.assertIsNotNone(runtime.session)
                await runtime.close()
                self.assertIsNone(runtime.session)

        asyncio.run(executar())
        self.assertEqual(inicializacoes, ["initialize"])
        self.assertEqual(encerramentos, ["sessao", "stdio"])

    def test_runtime_reinicia_sessao_apos_health_check_falhar(self):
        inicializacoes = []
        pings = []
        sessoes = []

        class FakeSession:
            async def __aenter__(self):
                return self

            async def __aexit__(self, *_args):
                return None

            async def initialize(self):
                inicializacoes.append(True)

            async def list_tools(self):
                return SimpleNamespace(tools=[])

            async def send_ping(self):
                pings.append(True)
                if len(pings) == 1:
                    raise RuntimeError("conexão perdida")

        @asynccontextmanager
        async def fake_stdio(_params):
            yield object(), object()

        def criar_sessao(*_args):
            sessao = FakeSession()
            sessoes.append(sessao)
            return sessao

        async def executar():
            runtime = main.MCPRuntime()
            with (
                patch.object(main, "stdio_client", fake_stdio),
                patch.object(main, "ClientSession", side_effect=criar_sessao),
                patch.object(main.os.path, "isfile", return_value=True),
            ):
                await runtime.start()
                self.assertFalse(await runtime.check())
                self.assertIsNotNone(runtime.startup_error)
                await runtime.start()
                self.assertEqual(len(inicializacoes), 2)
                self.assertIsNone(runtime.startup_error)
                self.assertIs(runtime.session, sessoes[1])
                await runtime.close()

        asyncio.run(executar())

    def test_agente_executa_chamadas_em_paralelo_e_sintetiza(self):
        chamada_a = main.types.Part.from_function_call(
            name="buscar_a",
            args={"termo": "A"},
        )
        chamada_b = main.types.Part.from_function_call(
            name="buscar_b",
            args={"termo": "B"},
        )

        class FakeModels:
            def __init__(self):
                self.turn = 0

            async def generate_content_stream(self, **_kwargs):
                self.turn += 1
                if self.turn == 1:
                    conteudo = main.types.Content(
                        role="model",
                        parts=[chamada_a, chamada_b],
                    )
                else:
                    conteudo = main.types.Content(
                        role="model",
                        parts=[main.types.Part.from_text(text="Conclusão baseada nos dados.")],
                    )
                yield SimpleNamespace(
                    candidates=[SimpleNamespace(content=conteudo)]
                )

        class FakeClient:
            def __init__(self):
                self.aio = SimpleNamespace(models=FakeModels(), aclose=self.aclose)
                self.closed = False

            async def aclose(self):
                return None

            def close(self):
                self.closed = True

        class FakeSession:
            def __init__(self):
                self.chamadas = []
                self.ativas = 0
                self.max_ativas = 0

            async def call_tool(self, nome, argumentos):
                self.chamadas.append((nome, argumentos))
                self.ativas += 1
                self.max_ativas = max(self.max_ativas, self.ativas)
                await asyncio.sleep(0.01)
                self.ativas -= 1
                return SimpleNamespace(
                    content=[SimpleNamespace(text=f"resultado {nome}")],
                    is_error=False,
                )

        class FakeRuntime:
            def __init__(self):
                self.session = FakeSession()
                self.tools = [
                    SimpleNamespace(
                        name="buscar_a",
                        description="Consulta A",
                        inputSchema={"type": "object", "properties": {}},
                    ),
                    SimpleNamespace(
                        name="buscar_b",
                        description="Consulta B",
                        inputSchema={"type": "object", "properties": {}},
                    ),
                ]
                self.vagas = asyncio.Semaphore(1)

            async def adquirir_vaga(self):
                await self.vagas.acquire()

            def liberar_vaga(self):
                self.vagas.release()

            async def start(self):
                return None

            async def call_tool(self, nome, argumentos):
                return await self.session.call_tool(nome, argumentos)

        async def executar():
            client = FakeClient()
            runtime = FakeRuntime()
            eventos = []

            async def registrar(evento):
                eventos.append(evento)

            with patch.object(main.genai, "Client", return_value=client):
                resultado = await main.processar_mcp_e_llm(
                    "Compare A e B",
                    [],
                    "PostgreSQL",
                    mcp_runtime=runtime,
                    on_event=registrar,
                )

            self.assertEqual(resultado[0], "Conclusão baseada nos dados.")
            self.assertTrue(resultado[2])
            self.assertEqual(len(runtime.session.chamadas), 2)
            self.assertEqual(runtime.session.max_ativas, 2)
            self.assertTrue(client.closed)
            self.assertEqual(
                [evento["type"] for evento in eventos].count("tool_start"),
                2,
            )

        with patch.dict("os.environ", {"GEMINI_API_KEY": "test-key"}):
            asyncio.run(executar())


class SqlValidationTests(unittest.TestCase):
    def test_sql_recusa_select_into(self):
        resultado = servidor_mcp.validar_e_executar_sql(
            "SELECT id INTO nova_tabela FROM tb_producao"
        )
        self.assertIn("operação de escrita", resultado)

    def test_sql_recusa_dml_em_cte(self):
        resultado = servidor_mcp.validar_e_executar_sql(
            "WITH removidos AS (DELETE FROM tb_producao RETURNING *) "
            "SELECT * FROM removidos"
        )
        self.assertIn("operação de escrita", resultado)

    def test_sql_recusa_schema_qualificado(self):
        resultado = servidor_mcp.validar_e_executar_sql(
            "SELECT id FROM public.tb_producao"
        )
        self.assertIn("schemas explícitos", resultado)

    def test_oracle_aceita_owner_dual_case_insensitive_e_rownum(self):
        consulta = servidor_mcp.validar_e_executar_sql(
            "SELECT ID FROM TI.TB_PRODUCAO WHERE ROWNUM <= 10",
            dialecto="oracle",
        )
        resultado_json = consulta.split("```json", 1)[1].split("```", 1)[0]
        resultado = json.loads(resultado_json)
        self.assertIn("ROWNUM", resultado["query_validada"])
        self.assertIn("TB_PRODUCAO", resultado["query_validada"])
        self.assertIn("FETCH FIRST 100 ROWS ONLY", resultado["query_validada"])

        resposta_dual = servidor_mcp.validar_e_executar_sql(
            "SELECT SYSDATE FROM DUAL",
            dialecto="oracle",
        )
        dual = json.loads(
            resposta_dual.split("```json", 1)[1].split("```", 1)[0]
        )
        self.assertIn("DUAL", dual["query_validada"])

    def test_oracle_recusa_tabela_desconhecida_e_dialeto_nao_suportado(self):
        resultado = servidor_mcp.validar_e_executar_sql(
            "SELECT id FROM tabela_inexistente",
            dialecto="oracle",
        )
        self.assertIn("não existe no ambiente", resultado)
        self.assertIn(
            "não é suportado",
            servidor_mcp.validar_e_executar_sql(
                "SELECT 1 FROM dual",
                dialecto="nao_existe",
            ),
        )


class FeedbackDatasetTests(unittest.TestCase):
    def test_feedback_negativo_exige_correcao_e_exporta_jsonl(self):
        argumentos = {
            "exchange_id": "exchange-1",
            "prompt": "Como consulto produção?",
            "resposta": "Resposta incorreta.",
            "avaliacao": "negative",
            "resposta_corrigida": "",
            "dialeto_sql": "Oracle",
            "modelo": "modelo-teste",
            "anexos": ["schema.sql"],
        }
        with self.assertRaisesRegex(ValueError, "resposta corrigida"):
            feedback.criar_registro_feedback(**argumentos)

        argumentos["resposta_corrigida"] = "Use SELECT com a tabela correta."
        registro = feedback.criar_registro_feedback(**argumentos)
        linha = feedback.exportar_jsonl([registro])
        recuperado = json.loads(linha)
        self.assertEqual(recuperado["resposta_alvo"], argumentos["resposta_corrigida"])
        self.assertTrue(recuperado["elegivel_para_treinamento_supervisionado"])

class RagTests(unittest.TestCase):
    def test_guard_usa_query_points_com_filtro_rbac(self):
        client = QdrantClient(location=":memory:")
        try:
            client.create_collection(
                collection_name="vetra_test",
                vectors_config=models.VectorParams(
                    size=2,
                    distance=models.Distance.COSINE,
                ),
            )
            client.upsert(
                collection_name="vetra_test",
                points=[
                    models.PointStruct(
                        id=str(uuid.uuid4()),
                        vector=[1.0, 0.0],
                        payload={
                            "departamento": "TI",
                            "nivel_acesso": 1,
                            "document": "Documento permitido",
                        },
                    ),
                    models.PointStruct(
                        id=str(uuid.uuid4()),
                        vector=[0.9, 0.1],
                        payload={
                            "departamento": "Financeiro",
                            "nivel_acesso": 3,
                            "document": "Documento restrito",
                        },
                    ),
                ],
                wait=True,
            )

            resultados = buscar_conhecimento_rag(
                client,
                "vetra_test",
                [1.0, 0.0],
                departamento_usuario="TI",
                nivel_acesso=1,
                limit=5,
            )

            self.assertEqual(len(resultados), 1)
            self.assertEqual(
                resultados[0].payload["departamento"],
                "TI",
            )
        finally:
            client.close()

    def test_indexacao_usa_upsert_ids_estaveis_e_remove_chunks_antigos(self):
        class FakeVector(list):
            def tolist(self):
                return list(self)

        class FakeEmbedding:
            def embed(self, documentos):
                return [FakeVector([1.0, 0.0]) for _ in documentos]

        class FakeClient:
            def __init__(self):
                self.upsert_args = None
                self.deleted_ids = None

            def upsert(self, **kwargs):
                self.upsert_args = kwargs

            def delete(self, **kwargs):
                self.deleted_ids = kwargs["points_selector"]

        client = FakeClient()
        servidor_mcp.CACHE_MEMORIA[
            'buscar_conhecimento_rag:"termo":{}'
        ] = {"timestamp": 1, "ttl": 300, "dados": "antigo"}
        with (
            patch.object(servidor_mcp, "obter_cliente_qdrant", return_value=client),
            patch.object(
                servidor_mcp,
                "obter_modelo_embedding",
                return_value=FakeEmbedding(),
            ),
            patch.object(
                servidor_mcp,
                "obter_ids_documento_qdrant",
                return_value=["id-antigo"],
            ),
        ):
            resultado = servidor_mcp.indexar_documento_com_chunking(
                "Texto corporativo",
                fonte="manual.txt",
                categoria="Operações",
            )

        self.assertIn("indexado com sucesso", resultado)
        self.assertEqual(len(client.upsert_args["points"]), 1)
        self.assertEqual(client.deleted_ids, ["id-antigo"])
        self.assertFalse(
            any(
                chave.startswith("buscar_conhecimento_rag:")
                for chave in servidor_mcp.CACHE_MEMORIA
            )
        )
